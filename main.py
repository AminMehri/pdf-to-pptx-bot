import os
from telegram import Update, InputFile
from telegram.ext import ApplicationBuilder, CommandHandler, MessageHandler, filters, ContextTypes
from pdf2image import convert_from_path
from pptx import Presentation
from pptx.util import Inches
from dotenv import load_dotenv

load_dotenv()


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("سلام! فایل PDF رو بفرست تا برات به پاورپوینت تبدیل کنم.")


async def handle_file(update: Update, context: ContextTypes.DEFAULT_TYPE):
    file = await update.message.document.get_file()
    file_name = update.message.document.file_name
    file_path = f"/tmp/{file_name}"
    await file.download_to_drive(file_path)

    # convert to pdf to images
    images = convert_from_path(file_path, dpi=150)

    # create pptx
    prs = Presentation()
    blank_slide = prs.slide_layouts[6]
    for i, img in enumerate(images):
        slide = prs.slides.add_slide(blank_slide)
        img_path = f"/tmp/page_{i + 1}.jpg"
        img.save(img_path, 'JPEG')
        slide.shapes.add_picture(img_path, Inches(0), Inches(0), width=prs.slide_width)

    pptx_path = f"/tmp/{file_name}.pptx"
    prs.save(pptx_path)

    with open(pptx_path, 'rb') as pptx_file:
        await update.message.reply_document(document=pptx_file, filename=f"{file_name}.pptx")

if __name__ == '__main__':
    import logging
    logging.basicConfig(level=logging.INFO)

    TOKEN = os.getenv("BOT_TOKEN")
    app = ApplicationBuilder().token(TOKEN).build()

    app.add_handler(CommandHandler("start", start))
    app.add_handler(MessageHandler(filters.Document.PDF, handle_file))

    print("Bot is running...")
    app.run_polling()